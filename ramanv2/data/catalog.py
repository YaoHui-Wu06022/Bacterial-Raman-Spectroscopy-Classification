"""数据集属、规范前缀和 Profile 的统一映射模型。"""

from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from ramanv2.common.publishing import create_temporary_path, publish_file

GRAM_PROFILE_NAMES = {"阴性": "GN", "阳性": "GP", "真菌": "FUNG"}
REQUIRED_COLUMNS = ("革兰氏", "属英文", "规范种简称")


@dataclass(frozen=True)
class DatasetMapping:
    """描述规范种简称到属以及属到 Profile 的唯一映射。"""

    genera: dict[str, tuple[str, ...]]
    datasets: dict[str, tuple[str, ...]]

    @classmethod
    def from_payload(cls, payload: object) -> "DatasetMapping":
        if not isinstance(payload, dict):
            raise ValueError("dataset_mapping 必须是 JSON 对象")
        raw_genera = payload.get("genera")
        raw_datasets = payload.get("datasets")
        if not isinstance(raw_genera, dict) or not isinstance(raw_datasets, dict):
            raise ValueError("dataset_mapping 必须同时包含 genera 和 datasets")

        genera: dict[str, tuple[str, ...]] = {}
        prefix_owner: dict[str, str] = {}
        for raw_name, raw_prefixes in raw_genera.items():
            genus = str(raw_name).strip()
            if not genus or not isinstance(raw_prefixes, list) or not raw_prefixes:
                raise ValueError("每个属必须有非空前缀列表")
            prefixes = tuple(str(prefix).strip().upper() for prefix in raw_prefixes if str(prefix).strip())
            if not prefixes:
                raise ValueError(f"属 {raw_name} 必须有非空前缀列表")
            if len(prefixes) != len(set(prefixes)):
                raise ValueError(f"属 {genus} 存在重复前缀")
            if genus in genera:
                raise ValueError(f"属 {genus} 重复定义")
            for prefix in prefixes:
                if prefix in prefix_owner:
                    raise ValueError(f"前缀 {prefix} 同时属于多个属")
                prefix_owner[prefix] = genus
            genera[genus] = prefixes

        datasets: dict[str, tuple[str, ...]] = {}
        for raw_name, raw_genera in raw_datasets.items():
            dataset_name = str(raw_name).strip()
            if not dataset_name or not isinstance(raw_genera, list) or not raw_genera:
                raise ValueError(f"数据集 {raw_name} 必须有非空属列表")
            names = tuple(str(name).strip() for name in raw_genera if str(name).strip())
            if len(names) != len(set(names)):
                raise ValueError(f"数据集 {dataset_name} 存在重复属")
            unknown = sorted(set(names) - set(genera))
            if unknown:
                raise ValueError(f"数据集 {dataset_name} 包含未定义属：{unknown}")
            if dataset_name in datasets:
                raise ValueError(f"数据集 {dataset_name} 重复定义")
            datasets[dataset_name] = names
        return cls(genera=genera, datasets=datasets)

    @classmethod
    def from_json(cls, path: Path | str) -> "DatasetMapping":
        return cls.from_payload(json.loads(Path(path).read_text(encoding="utf-8")))

    @classmethod
    def from_workbook(cls, workbook_path: Path | str) -> "DatasetMapping":
        path = Path(workbook_path).resolve()
        if not path.is_file():
            raise FileNotFoundError(f"病原菌分类表不存在：{path}")
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            rows = workbook.active.iter_rows(values_only=True)
            headers = next(rows, None)
            if headers is None:
                raise ValueError("病原菌分类表为空")
            positions = {str(value).strip(): index for index, value in enumerate(headers) if value is not None}
            missing = [name for name in REQUIRED_COLUMNS if name not in positions]
            if missing:
                raise ValueError(f"病原菌分类表缺少列：{missing}")

            genera_lists: dict[str, list[str]] = {}
            genus_grams: dict[str, str] = {}
            prefix_owner: dict[str, str] = {}
            profile_genera: dict[str, set[str]] = {name: set() for name in ("GN", "GP", "FUNG")}
            for row in rows:
                gram = str(row[positions["革兰氏"]] or "").strip()
                genus = str(row[positions["属英文"]] or "").strip()
                prefix = str(row[positions["规范种简称"]] or "").strip().upper()
                if not gram and not genus and not prefix:
                    continue
                if gram not in GRAM_PROFILE_NAMES or not genus or not prefix:
                    raise ValueError(f"病原菌分类表存在无效记录：{row}")
                previous_gram = genus_grams.setdefault(genus, gram)
                if previous_gram != gram:
                    raise ValueError(f"属 {genus} 同时出现在多个分类：{previous_gram}、{gram}")
                owner = prefix_owner.get(prefix)
                if owner is not None and owner != genus:
                    raise ValueError(f"规范种简称 {prefix} 同时属于 {owner} 和 {genus}")
                prefix_owner[prefix] = genus
                if prefix not in genera_lists.setdefault(genus, []):
                    genera_lists[genus].append(prefix)
                profile_genera[GRAM_PROFILE_NAMES[gram]].add(genus)
        finally:
            workbook.close()

        genera = {name: tuple(prefixes) for name, prefixes in sorted(genera_lists.items())}
        datasets = {
            "GN": tuple(sorted(profile_genera["GN"])),
            "GP": tuple(sorted(profile_genera["GP"])),
            "FUNG": tuple(sorted(profile_genera["FUNG"])),
            "MICRO": tuple(sorted(genera)),
        }
        return cls(genera=genera, datasets=datasets)

    def subset(self, profile_name: str) -> "DatasetMapping":
        """生成只包含一个 Profile 的映射，供单 Profile 构建使用。"""
        if profile_name not in self.datasets:
            raise ValueError(f"未知 Profile：{profile_name}")
        return DatasetMapping(genera=dict(self.genera), datasets={profile_name: self.datasets[profile_name]})

    def to_payload(self) -> dict[str, object]:
        return {
            "version": 1,
            "datasets": {name: list(genera) for name, genera in self.datasets.items()},
            "genera": {name: list(prefixes) for name, prefixes in self.genera.items()},
        }


def load_dataset_mapping(path: Path | str) -> DatasetMapping:
    """读取 UTF-8 JSON 映射文件。"""
    return DatasetMapping.from_json(path)


def ensure_dataset_mapping(workbook_path: Path | str, mapping_path: Path | str) -> DatasetMapping:
    """按 Excel 修改时间生成或复用唯一的本地 JSON 映射。"""
    workbook = Path(workbook_path).resolve()
    target = Path(mapping_path).resolve()
    if target.is_file() and (not workbook.is_file() or target.stat().st_mtime_ns >= workbook.stat().st_mtime_ns):
        try:
            return load_dataset_mapping(target)
        except (OSError, ValueError, json.JSONDecodeError):
            # 配置可能来自旧页面或被中断写入；根据当前 Excel 原子重建。
            if not workbook.is_file():
                raise

    mapping = DatasetMapping.from_workbook(workbook)
    temporary = create_temporary_path(target)
    try:
        temporary.write_text(json.dumps(mapping.to_payload(), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        publish_file(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)
    return mapping


__all__ = [
    "DatasetMapping",
    "ensure_dataset_mapping",
    "load_dataset_mapping",
]
